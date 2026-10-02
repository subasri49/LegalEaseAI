"""Formatting utilities: sanitize text and export to DOCX, PDF and an HTML preview."""
import html
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos

import config

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " ", "\u2022": "-",
}


# ----------------------------------------------------------------------------- helpers
def sanitize_text(text: str) -> str:
    """Remove typographic characters and markdown symbols for clean formatting."""
    if not text:
        return ""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("\r\n", "\n").replace("`", "")
    text = re.sub(r"(?m)^#{1,6}\s*", "", text)           # markdown headings
    text = re.sub(r"(?m)^\s*\*\s+", "- ", text)           # "* item" -> "- item"
    text = re.sub(r"\*+", "", text)                       # bold / stray asterisks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_terms(terms: str):
    return [t.strip() for t in (terms or "").replace("\n", ";").split(";") if t.strip()]


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if re.match(r"^(ARTICLE|SECTION)\s+\w+", s, re.I):
        return True
    if re.match(r"^\d+(\.\d+)*[\.\)]\s+\S", s) and not s.endswith("."):
        return True
    if s.isupper() and len(s) > 3:
        return True
    return s.endswith(":") and len(s.split()) <= 7


def is_bullet(line: str) -> bool:
    return bool(re.match(r"^\s*(-|[a-z]\)|\([a-z]\))\s+", line))


def parse_blocks(text: str):
    """Returns [(kind, content)] where kind is title | heading | bullet | para."""
    blocks = []
    first = True
    for line in sanitize_text(text).split("\n"):
        if not line.strip():
            continue
        if first:
            blocks.append(("title", line.strip()))
            first = False
        elif is_heading(line):
            blocks.append(("heading", line.strip()))
        elif is_bullet(line):
            blocks.append(("bullet", re.sub(r"^\s*-\s+", "", line).strip()))
        else:
            blocks.append(("para", line.strip()))
    return blocks


def _logo(path=None):
    p = Path(path or config.LOGO_PATH)
    return p if p.exists() else None


# ----------------------------------------------------------------------------- DOCX
def _add_page_number(paragraph):
    run = paragraph.add_run(" | Page ")
    run.font.size = Pt(8)
    run2 = paragraph.add_run()
    run2.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run2._r.append(el)


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    doc.core_properties.title = doc_type
    doc.core_properties.author = "LegalEase"

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(config.FOOTER_TEXT)
        r.italic = True
        r.font.size = Pt(8)
        _add_page_number(fp)

    logo = _logo()
    if logo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), width=Inches(2.0))

    for kind, content in parse_blocks(text):
        if kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(content)
            r.bold = True
            r.font.size = Pt(16)
        elif kind == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
            p.add_run(content).bold = True
        elif kind == "bullet":
            doc.add_paragraph(content, style="List Bullet")
        else:
            p = doc.add_paragraph(content)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    items = split_terms(terms)
    if items:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"
        for c in table.rows[0].cells:
            for run in c.paragraphs[0].runs:
                run.bold = True
        for i, item in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = sanitize_text(item)
        for row in table.rows:
            row.cells[0].width = Inches(0.6)
            row.cells[1].width = Inches(5.9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- PDF
def _latin(s: str) -> str:
    """Core PDF fonts only support Latin-1."""
    return s.encode("latin-1", "replace").decode("latin-1")


class _LegalPDF(FPDF):
    def __init__(self, doc_type):
        super().__init__(format="A4")
        self.doc_type = doc_type
        self.set_margins(20, 20, 20)
        self.set_auto_page_break(True, margin=22)
        self.alias_nb_pages()

    def header(self):
        logo = _logo()
        if logo:
            self.image(str(logo), x=(self.w - 42) / 2, y=8, w=42)
        self.set_y(26)
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(90, 90, 90)
        self.cell(0, 5, _latin(self.doc_type), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_text_color(0, 0, 0)
        self.ln(4)

    def footer(self):
        self.set_y(-16)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 5, _latin(f"{config.FOOTER_TEXT} | Page {self.page_no()}/{{nb}}"), align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    pdf = _LegalPDF(doc_type)
    pdf.set_title(_latin(doc_type))
    pdf.set_author("LegalEase")
    pdf.add_page()

    def write(txt, font, style, size, h, align="L", gap=0):
        pdf.set_font(font, style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, h, _latin(txt), align=align, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if gap:
            pdf.ln(gap)

    for kind, content in parse_blocks(text):
        if kind == "title":
            write(content, "Helvetica", "B", 15, 8, "C", 4)
        elif kind == "heading":
            pdf.ln(2)
            write(content, "Helvetica", "B", 11, 6, "L", 1)
        elif kind == "bullet":
            write("-  " + content, "Helvetica", "", 10.5, 5.5, "L", 1)
        else:
            write(content, "Helvetica", "", 10.5, 5.5, "L", 2)

    items = split_terms(terms)
    if items:
        pdf.ln(4)
        write("Summary of Key Terms", "Helvetica", "B", 11, 6, "L", 1)
        for item in items:
            write("-  " + sanitize_text(item), "Helvetica", "", 10.5, 5.5, "L", 1)

    return bytes(pdf.output())


# ----------------------------------------------------------------------------- HTML preview
def format_html_preview(text: str) -> str:
    out = []
    for kind, content in parse_blocks(text):
        c = html.escape(content)
        if kind == "title":
            out.append(f"<h3 style='text-align:center;margin:0 0 14px'>{c}</h3>")
        elif kind == "heading":
            out.append(f"<h4 style='margin:16px 0 4px;color:#9ecbff'>{c}</h4>")
        elif kind == "bullet":
            out.append(f"<div style='margin:2px 0 2px 18px'>&bull; {c}</div>")
        else:
            out.append(f"<p style='margin:6px 0;line-height:1.55'>{c}</p>")
    return "".join(out)
