"""Run with:  pytest -v   (no real Gemini key needed - the AI call is mocked)."""
import io
import zipfile

import pytest
from fastapi.testclient import TestClient

import config
from ai_core import generator
from legalEaseAPI import routes
from legalEaseAPI.main import app

SAMPLE = """## **FREELANCE WORK CONTRACT**
Agreement made this \u201c15th\u201d day of April, 2025
Between:
Jane Doe (the "Service Provider") and TechNova Inc. (the "Client")
1. Services:
The Service Provider shall deliver the work by May 15, 2025.
2. Payment:
- Payment within 7 days of invoice
- Late fees may apply
"""


class FakeGenerator:
    last_model_used = "fake-model"

    def generate_document(self, *args):
        return SAMPLE


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(routes, "_generator", FakeGenerator())
    return TestClient(app)


PAYLOAD = {"document_type": "Freelance Work Contract", "parties": "Jane Doe; TechNova Inc.",
           "terms": "Work by May 15; Pay in 7 days", "dates": "April 15, 2025"}


def test_root_and_health(client):
    assert client.get("/").status_code == 200
    assert client.get("/health").json() == {"status": "ok"}


def test_generate_and_history_crud(client):
    r = client.post("/generate", json=PAYLOAD)
    assert r.status_code == 200
    doc_id = r.json()["id"]
    assert client.get("/history").json()["items"][0]["id"] == doc_id
    assert client.put(f"/history/{doc_id}", json={"content": "edited"}).status_code == 200
    assert client.get(f"/history/{doc_id}").json()["content"] == "edited"
    assert client.delete(f"/history/{doc_id}").status_code == 200
    assert client.get(f"/history/{doc_id}").status_code == 404


def test_validation_error(client):
    assert client.post("/generate", json={"document_type": "x"}).status_code == 422


def test_missing_api_key_returns_503(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(config, "GEMINI_API_KEY", "")
    monkeypatch.setattr(routes, "_generator", None)
    assert TestClient(app).post("/generate", json=PAYLOAD).status_code == 503


def test_sanitize_removes_markdown_and_smart_quotes():
    out = generator.sanitize_text(SAMPLE)
    assert "#" not in out and "*" not in out and "\u201c" not in out


def test_docx_has_terms_table():
    data = generator.format_docx(SAMPLE, "Freelance Work Contract", "Work by May 15; Pay in 7 days")
    assert zipfile.is_zipfile(io.BytesIO(data))
    from docx import Document
    d = Document(io.BytesIO(data))
    assert len(d.tables) == 1 and len(d.tables[0].rows) == 3


def test_pdf_is_valid():
    data = generator.format_pdf(SAMPLE * 30, "Freelance Work Contract", "A; B")
    assert data.startswith(b"%PDF")


def test_html_preview_escapes():
    assert "&lt;script&gt;" in generator.format_html_preview("TITLE\n<script>")
